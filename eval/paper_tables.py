"""Build paper-ready experiment tables from evaluation outputs and annotations.

This script is intentionally conservative: it only reports metrics that can be
traced to result files or explicit annotation files. Missing annotations are
rendered as TBD instead of being inferred.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path
from statistics import mean, pstdev
from typing import Any


ROOT = Path(__file__).resolve().parents[1]
VARIANTS = ("full", "wo_routing", "wo_bm25", "wo_schema", "wo_images", "wo_multi_agent")
PAIR_SUBSETS = {
    "wo_routing": "routing",
    "wo_bm25": "retrieval",
    "wo_schema": "sql",
    "wo_images": "deep",
    "wo_multi_agent": "deep",
}
QUALITY_COLUMNS = ("correctness", "completeness", "traceability", "actionability")


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def load_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def load_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def pct(value: float | None) -> str:
    return "TBD" if value is None else f"{value * 100:.1f}%"


def num(value: float | int | None, digits: int = 2) -> str:
    if value is None:
        return "TBD"
    if isinstance(value, int):
        return str(value)
    return f"{value:.{digits}f}"


def md_cell(value: Any) -> str:
    text = "TBD" if value in (None, "") else str(value)
    return text.replace("|", r"\|").replace("\n", "<br>")


def rate(rows: list[dict[str, Any]], predicate) -> float | None:
    if not rows:
        return None
    return sum(1 for row in rows if predicate(row)) / len(rows)


def avg(values: list[Any]) -> float | None:
    clean = [float(v) for v in values if isinstance(v, (int, float))]
    return mean(clean) if clean else None


def mean_sd(values: list[Any]) -> str:
    clean = [float(v) for v in values if isinstance(v, (int, float))]
    if not clean:
        return "TBD"
    if len(clean) == 1:
        return f"{clean[0]:.2f}"
    return f"{mean(clean):.2f} +/- {pstdev(clean):.2f}"


def table_names(sqls: list[str]) -> set[str]:
    names: set[str] = set()
    for sql in sqls:
        names.update(re.findall(r"(?i)\b(?:from|join)\s+[`\"]?([\w.]+)", str(sql)))
    return names


def sql_ready(row: dict[str, Any]) -> bool:
    sqls = row.get("executed_sqls") or []
    terms = [str(x).lower() for x in row.get("gold_sql_terms") or []]
    blob = " ".join(sqls).lower()
    table_hit = bool(set(row.get("gold_tables") or []) & table_names(sqls))
    term_hit = bool(terms) and all(term in blob for term in terms)
    return bool(sqls) and table_hit and term_hit


def first_relevant_rank(row: dict[str, Any]) -> int | None:
    gold_ids = {str(x).lower() for x in row.get("gold_evidence_ids") or []}
    gold_keywords = [str(x).lower() for x in row.get("gold_keywords") or []]
    for index, chunk in enumerate(row.get("evidence_chunks") or [], 1):
        blob = json.dumps(chunk, ensure_ascii=False).lower()
        chunk_id = str(chunk.get("id") or chunk.get("chunk_id") or chunk.get("source_id") or "").lower()
        if chunk_id and chunk_id in gold_ids:
            return index
        if gold_keywords and any(term in blob for term in gold_keywords):
            return index
    return None


def recall_at(rows: list[dict[str, Any]], k: int) -> float | None:
    eligible = [
        row
        for row in rows
        if row.get("gold_evidence_ids") or row.get("gold_keywords")
    ]
    return rate(eligible, lambda row: (first_relevant_rank(row) or math.inf) <= k)


def mrr(rows: list[dict[str, Any]]) -> float | None:
    eligible = [
        row
        for row in rows
        if row.get("gold_evidence_ids") or row.get("gold_keywords")
    ]
    values = []
    for row in eligible:
        rank = first_relevant_rank(row)
        values.append(0.0 if rank is None else 1.0 / rank)
    return avg(values)


def visual_recall(rows: list[dict[str, Any]]) -> float | None:
    eligible = [row for row in rows if row.get("requires_image")]
    return rate(
        eligible,
        lambda row: bool(row.get("kb_images"))
        or int((row.get("ablation_trace") or {}).get("injected_image_count") or 0) > 0,
    )


def result_metrics(rows: list[dict[str, Any]]) -> dict[str, Any]:
    done = [row for row in rows if row.get("status") == "done"]
    sql_rows = [row for row in done if row.get("subset") == "sql"]
    routing_rows = [row for row in done if row.get("gold_action")]
    retrieval_rows = [row for row in done if row.get("subset") == "retrieval"]
    deep_rows = [row for row in done if row.get("subset") in {"deep", "end_to_end"}]
    return {
        "task_count": len(rows),
        "done": len(done),
        "completion": len(done) / len(rows) if rows else None,
        "action_accuracy": rate(
            routing_rows, lambda row: row.get("observed_action") == row.get("gold_action")
        ),
        "sql_ready": rate(sql_rows, sql_ready),
        "retrieval_recall3": recall_at(retrieval_rows, 3),
        "retrieval_recall5": recall_at(retrieval_rows, 5),
        "retrieval_mrr": mrr(retrieval_rows),
        "visual_recall": visual_recall(deep_rows + retrieval_rows),
        "latency": avg([row.get("answer_elapsed_seconds") for row in done]),
        "latency_mean_sd": mean_sd([row.get("answer_elapsed_seconds") for row in done]),
    }


def load_variant_results(root: Path, experiment_id: str) -> dict[str, list[dict[str, Any]]]:
    exp = root / experiment_id
    return {variant: load_jsonl(exp / variant / "results.jsonl") for variant in VARIANTS}


def load_manual_scores(path: Path) -> dict[str, dict[str, float]]:
    scores: dict[str, dict[str, float]] = {}
    for row in load_csv(path):
        key = row.get("task_id") or row.get("id")
        variant = row.get("variant", "full")
        if not key:
            continue
        values: dict[str, float] = {}
        for col in (*QUALITY_COLUMNS, "root_cause_correct", "unsupported_claim_rate"):
            raw = row.get(col, "")
            if raw == "":
                continue
            try:
                values[col] = float(raw)
            except ValueError:
                continue
        if values:
            scores[f"{variant}:{key}"] = values
    return scores


def quality_summary(rows: list[dict[str, Any]], manual_scores: dict[str, dict[str, float]], variant: str) -> dict[str, Any]:
    keys = [f"{variant}:{row.get('id')}" for row in rows]
    matched = [manual_scores[key] for key in keys if key in manual_scores]
    summary: dict[str, Any] = {"scored": len(matched)}
    for col in QUALITY_COLUMNS:
        summary[col] = avg([row.get(col) for row in matched])
    summary["quality"] = avg([
        mean([value for metric, value in row.items() if metric in QUALITY_COLUMNS])
        for row in matched
        if any(metric in row for metric in QUALITY_COLUMNS)
    ])
    summary["root_cause_correct"] = avg([row.get("root_cause_correct") for row in matched])
    summary["unsupported_claim_rate"] = avg([row.get("unsupported_claim_rate") for row in matched])
    return summary


def render_table(headers: list[str], rows: list[list[Any]], aligns: list[str] | None = None) -> str:
    aligns = aligns or ["---"] * len(headers)
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(aligns) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(md_cell(value) for value in row) + " |")
    return "\n".join(lines) + "\n"


def table_i_dataset_stats(dataset_stats: dict[str, Any], all_rows: list[dict[str, Any]]) -> str:
    subset_counts = Counter(row.get("subset", "unknown") for row in all_rows)
    rows = [
        ["Structured data sources", dataset_stats.get("structured_sources")],
        ["Operating-data time span", dataset_stats.get("time_span")],
        ["Alarm events", dataset_stats.get("alarm_event_count")],
        ["Telemetry fields", dataset_stats.get("telemetry_field_count")],
        ["Diagnostic result records", dataset_stats.get("diagnostic_result_count")],
        ["Knowledge-base documents", dataset_stats.get("kb_document_count")],
        ["Text chunks", dataset_stats.get("kb_chunk_count")],
        ["Linked images/figures", dataset_stats.get("kb_image_count")],
        ["Historical cases", dataset_stats.get("historical_case_count")],
        ["Evaluation tasks", len(all_rows) if all_rows else dataset_stats.get("evaluation_task_count")],
        ["Task split", ", ".join(f"{k}={v}" for k, v in sorted(subset_counts.items())) or dataset_stats.get("task_split")],
    ]
    return "## Table I. Dataset and Task-Suite Statistics\n\n" + render_table(
        ["Item", "Value"], rows
    )


def table_ii_main(results: dict[str, list[dict[str, Any]]], manual_scores: dict[str, dict[str, float]]) -> str:
    rows = []
    for label, variant in [
        ("Pure LLM", "pure_llm"),
        ("RAG-only", "rag_only"),
        ("Tool-only / Text-to-SQL-only", "tool_only"),
        ("Proposed full system", "full"),
    ]:
        variant_rows = results.get(variant, [])
        metrics = result_metrics(variant_rows)
        quality = quality_summary(variant_rows, manual_scores, variant)
        rows.append([
            label,
            metrics["task_count"] or "TBD",
            num(quality["quality"]),
            pct(metrics["retrieval_recall5"]),
            pct(metrics["sql_ready"]),
            pct(metrics["action_accuracy"]),
            num(metrics["latency"]),
            "TBD",
        ])
    return "## Table II. Main Comparison with Baselines\n\n" + render_table(
        [
            "System",
            "Tasks",
            "Quality",
            "Retrieval Hit",
            "SQL Success",
            "Action Acc.",
            "Latency (s)",
            "Token Cost",
        ],
        rows,
        ["---", "---:", "---:", "---:", "---:", "---:", "---:", "---:"],
    )


def table_iii_retrieval(results: dict[str, list[dict[str, Any]]]) -> str:
    rows = []
    variants = [
        ("Dense retrieval", "dense"),
        ("BM25", "bm25"),
        ("Hybrid text retrieval", "full"),
        ("Hybrid text-image retrieval", "full"),
        ("w/o BM25", "wo_bm25"),
        ("w/o Images", "wo_images"),
    ]
    for label, variant in variants:
        variant_rows = [
            row
            for row in results.get(variant, [])
            if row.get("subset") in {"retrieval", "deep"}
        ]
        metrics = result_metrics(variant_rows)
        rows.append([
            label,
            metrics["task_count"] or "TBD",
            pct(metrics["retrieval_recall3"]),
            pct(metrics["retrieval_recall5"]),
            num(metrics["retrieval_mrr"]),
            pct(metrics["visual_recall"]),
            num(metrics["latency"]),
        ])
    return "## Table III. Retrieval Quality Comparison\n\n" + render_table(
        ["Retriever", "Tasks", "Recall@3", "Recall@5", "MRR", "Visual Recall", "Latency (s)"],
        rows,
        ["---", "---:", "---:", "---:", "---:", "---:", "---:"],
    )


def table_iv_routing(results: dict[str, list[dict[str, Any]]], manual_scores: dict[str, dict[str, float]]) -> str:
    rows = []
    for label, variant in [("Full", "full"), ("w/o Routing", "wo_routing")]:
        variant_rows = [row for row in results.get(variant, []) if row.get("subset") == "routing"]
        metrics = result_metrics(variant_rows)
        quality = quality_summary(variant_rows, manual_scores, variant)
        deep_rate = rate(variant_rows, lambda row: row.get("observed_action") == "DEEP")
        rows.append([
            label,
            metrics["task_count"],
            pct(metrics["action_accuracy"]),
            pct(deep_rate),
            num(quality["quality"]),
            metrics["latency_mean_sd"],
        ])
    return "## Table IV. Ablation of Adaptive Routing\n\n" + render_table(
        ["Variant", "Tasks", "Action Acc.", "Deep-Path Rate", "Quality", "Latency (s)"],
        rows,
        ["---", "---:", "---:", "---:", "---:", "---:"],
    )


def table_v_sql(results: dict[str, list[dict[str, Any]]], sql_results: dict[str, list[dict[str, Any]]]) -> str:
    source = sql_results if sql_results.get("full") else results
    rows = []
    for label, variant in [("Full", "full"), ("w/o Schema", "wo_schema")]:
        variant_rows = [row for row in source.get(variant, []) if row.get("subset") == "sql"]
        sql_generation = rate(variant_rows, lambda row: bool(row.get("executed_sqls")))
        table_hit = rate(
            variant_rows,
            lambda row: bool(set(row.get("gold_tables") or []) & table_names(row.get("executed_sqls") or [])),
        )
        term_hit = rate(
            variant_rows,
            lambda row: bool(row.get("gold_sql_terms"))
            and all(
                str(term).lower() in " ".join(row.get("executed_sqls") or []).lower()
                for term in row.get("gold_sql_terms") or []
            ),
        )
        rows.append([
            label,
            len(variant_rows),
            pct(sql_generation),
            pct(table_hit),
            pct(term_hit),
            pct(rate(variant_rows, sql_ready)),
            mean_sd([row.get("answer_elapsed_seconds") for row in variant_rows]),
        ])
    return "## Table V. Ablation of Schema-Constrained SQL Generation\n\n" + render_table(
        ["Variant", "Tasks", "SQL Generation", "Table Hit", "Condition Hit", "SQL-Ready", "Latency (s)"],
        rows,
        ["---", "---:", "---:", "---:", "---:", "---:", "---:"],
    )


def table_vi_multi_agent(results: dict[str, list[dict[str, Any]]], manual_scores: dict[str, dict[str, float]]) -> str:
    rows = []
    for label, variant in [("Full", "full"), ("w/o Multi-Agent", "wo_multi_agent")]:
        variant_rows = [row for row in results.get(variant, []) if row.get("subset") == "deep"]
        quality = quality_summary(variant_rows, manual_scores, variant)
        rows.append([
            label,
            len(variant_rows),
            num(quality["quality"]),
            num(quality["traceability"]),
            num(quality["completeness"]),
            num(quality["actionability"]),
            mean_sd([row.get("answer_elapsed_seconds") for row in variant_rows]),
        ])
    return "## Table VI. Ablation of Multi-Agent Diagnosis\n\n" + render_table(
        ["Variant", "Tasks", "Quality", "Traceability", "Completeness", "Actionability", "Latency (s)"],
        rows,
        ["---", "---:", "---:", "---:", "---:", "---:", "---:"],
    )


def table_vii_end_to_end(results: dict[str, list[dict[str, Any]]], manual_scores: dict[str, dict[str, float]]) -> str:
    rows = []
    for label, variant in [("Proposed full system", "full"), ("RAG-only", "rag_only"), ("Pure LLM", "pure_llm")]:
        variant_rows = [row for row in results.get(variant, []) if row.get("subset") == "end_to_end"]
        quality = quality_summary(variant_rows, manual_scores, variant)
        rows.append([
            label,
            len(variant_rows) if variant_rows else "TBD",
            pct(quality["root_cause_correct"]),
            num(quality["traceability"]),
            num(quality["completeness"]),
            num(quality["actionability"]),
            pct(quality["unsupported_claim_rate"]),
        ])
    return "## Table VII. End-to-End Diagnostic Case Evaluation\n\n" + render_table(
        [
            "System",
            "Cases",
            "Root-Cause Correct",
            "Traceability",
            "Completeness",
            "Actionability",
            "Unsupported Claims",
        ],
        rows,
        ["---", "---:", "---:", "---:", "---:", "---:", "---:"],
    )


def write_docx(markdown: str, out_path: Path) -> None:
    try:
        from docx import Document
    except Exception:
        return
    doc = Document()
    for block in markdown.splitlines():
        if block.startswith("# "):
            doc.add_heading(block[2:], level=1)
        elif block.startswith("## "):
            doc.add_heading(block[3:], level=2)
        elif block.startswith("|"):
            # Keep Markdown tables as monospaced text for traceable copy-paste.
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(block)
            run.font.name = "Consolas"
        elif block.strip():
            doc.add_paragraph(block)
        else:
            doc.add_paragraph("")
    doc.save(out_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=ROOT / "eval_outputs" / "v2")
    parser.add_argument("--experiment-id", default="paper_v2")
    parser.add_argument("--sql-experiment-id", default="paper_v2_sql20")
    parser.add_argument(
        "--dataset-stats",
        type=Path,
        default=ROOT / "eval" / "datasets" / "paper_dataset_stats.template.json",
    )
    parser.add_argument(
        "--manual-scores",
        type=Path,
        default=ROOT / "eval" / "datasets" / "paper_manual_scores.template.csv",
    )
    parser.add_argument("--out", type=Path, default=ROOT / "eval_outputs" / "v2" / "paper_tables")
    args = parser.parse_args()

    args.out.mkdir(parents=True, exist_ok=True)
    results = load_variant_results(args.root, args.experiment_id)
    sql_results = load_variant_results(args.root, args.sql_experiment_id)
    dataset_stats = load_json(args.dataset_stats)
    manual_scores = load_manual_scores(args.manual_scores)
    all_rows = results.get("full") or []

    tables = [
        "# Paper Experiment Result Tables",
        "",
        table_i_dataset_stats(dataset_stats, all_rows),
        table_ii_main(results, manual_scores),
        table_iii_retrieval(results),
        table_iv_routing(results, manual_scores),
        table_v_sql(results, sql_results),
        table_vi_multi_agent(results, manual_scores),
        table_vii_end_to_end(results, manual_scores),
    ]
    combined = "\n".join(tables)
    (args.out / "paper_experiment_results.md").write_text(combined, encoding="utf-8")
    write_docx(combined, args.out / "paper_experiment_results.docx")

    for name, content in {
        "table_i_dataset_stats.md": table_i_dataset_stats(dataset_stats, all_rows),
        "table_ii_main_comparison.md": table_ii_main(results, manual_scores),
        "table_iii_retrieval_quality.md": table_iii_retrieval(results),
        "table_iv_routing_ablation.md": table_iv_routing(results, manual_scores),
        "table_v_sql_ablation.md": table_v_sql(results, sql_results),
        "table_vi_multi_agent_ablation.md": table_vi_multi_agent(results, manual_scores),
        "table_vii_end_to_end.md": table_vii_end_to_end(results, manual_scores),
    }.items():
        (args.out / name).write_text(content, encoding="utf-8")

    summary = {
        "experiment_id": args.experiment_id,
        "sql_experiment_id": args.sql_experiment_id,
        "output_dir": str(args.out),
        "full_task_count": len(results.get("full", [])),
        "manual_score_count": len(manual_scores),
        "variants_with_results": [variant for variant, rows in results.items() if rows],
    }
    (args.out / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
